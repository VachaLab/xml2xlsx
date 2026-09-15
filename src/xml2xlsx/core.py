# Released under GPL3 License.
# Copyright (c) 2026 Ladislav Bartos and Robert Vacha Lab

"""
Convert a microplate reader's XML export into an XLSX workbook.

The input file is SpreadsheetML 2003: XML that describes a spreadsheet,
one row and cell at a time. It can hold several plates below each other
on one worksheet, each starting with a "Plate" header row. The script
finds them and copies each one to its own sheet of a new .xlsx file,
keeping the cells and their formatting as they were.
"""

import re
from dataclasses import dataclass
from pathlib import Path
from xml.etree.ElementTree import Element, ParseError, parse

from openpyxl.cell.cell import Cell as XlsxCell
from openpyxl.styles.alignment import Alignment
from openpyxl.styles.fonts import Font
from openpyxl.utils.cell import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# every tag and attribute in these files lives in this namespace
SS = "urn:schemas-microsoft-com:office:spreadsheet"

# matches plate header row
PLATE_HEADER = re.compile(r"^plate\b(?:\s*\(\s*\d+\s+of\s+\d+\s*\))?$", re.IGNORECASE)

# matches experiment row
EXPERIMENT = re.compile(
    r"^experiment\b(?:\s*\(\s*\d+\s+of\s+\d+\s*\))?$", re.IGNORECASE
)

# matches well-plate row label
ROW_LABEL = re.compile(r"^[A-Z]{1,2}$", re.IGNORECASE)

# characters Excel forbids in a sheet name
BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

# placeholder read time for "never read"
NEVER_READ = "01/01/0001"

# format codes the XML names but Excel does not understand
NUMBER_FORMATS = {"general date": "m/d/yyyy h:mm:ss"}

# the XML gives column widths in points, openpyxl wants them in characters
POINTS_PER_CHARACTER = 5.25

# columns of the report table, in order
REPORT_HEADERS = ("sheet", "plate", "experiment", "read time", "source", "status")


def tag(name: str) -> str:
    """Prefix a tag name with the SpreadsheetML namespace."""
    return f"{{{SS}}}{name}"


def attr(element: Element, name: str) -> str | None:
    """Read one of the element's `ss:` attributes, e.g. Index or StyleID."""
    return element.get(tag(name))


def children(element: Element, name: str) -> list[Element]:
    """The element's direct children with that name, `[]` if it has none."""
    return element.findall(tag(name))


def index_of(raw: str | None, following: int) -> int:
    """
    Work out where a row or cell sits.

    Args:
        raw: The `ss:Index` attribute, if the item has one.
        following: Where the item lands if it does not.
    """
    return int(raw) if raw is not None and raw.isdigit() else following


def normalize(text: str) -> str:
    """Convert a label into a comparable form: one space, no trailing colon, lowercase."""
    return re.sub(r"\s+", " ", text).strip().rstrip(":").strip().casefold()


@dataclass(frozen=True)
class Style:
    """
    A `<Style>` element translated into the objects openpyxl wants.

    The XML keeps formatting in one table of named styles and each cell
    points at one by ID.

    openpyxl does not have any such table. Here, a Font or an Alignment object
    is assigned to the cell itself. So each style is translated once, here,
    and then handed to every cell that uses it.

    Only the styles these files actually use are covered.
    """

    font: Font | None
    alignment: Alignment | None
    number_format: str | None

    def apply_to(self, cell: XlsxCell) -> None:
        """Put this style on a cell of the workbook being written."""
        if self.font is not None:
            cell.font = self.font
        if self.alignment is not None:
            cell.alignment = self.alignment
        if self.number_format is not None:
            cell.number_format = self.number_format


@dataclass(frozen=True)
class Cell:
    """
    One `<Cell>` of the source file.

    Attributes:
        row: 1-based row, as the spreadsheet counts them.
        col: 1-based column, likewise.
        value: int, float or str, or None if the cell is empty.
        style_id: Name of the style to use, or None for the default one.
        merge_across: Extra columns this cell is merged over.
            0 = ordinary cell, 1 = two columns, 2 = three columns etc.
    """

    row: int
    col: int
    value: object
    style_id: str | None
    merge_across: int


@dataclass(frozen=True)
class Sheet:
    """
    One worksheet of the source file, as a row number -> cells mapping.

    Both levels are sparse: a row missing from `rows` held nothing at
    all, and the cells within a row need not be adjacent. They are in
    left-to-right order. Widths and heights hold only what the file
    bothered to state, keyed by column and row number.

    Attributes:
        rows: The worksheet's cells, grouped by row.
        widths: Column widths, in points.
        heights: Row heights, in points.
        last_row: Number of the last row that held anything.
    """

    rows: dict[int, list[Cell]]
    widths: dict[int, float]
    heights: dict[int, float]
    last_row: int

    def text(self, cell: Cell | None) -> str:
        """A cell's value as a stripped string. Empty if it has none."""
        return "" if cell is None or cell.value is None else str(cell.value).strip()

    def only_value(self, row: int) -> str | None:
        """
        The row's text if it holds exactly one value, else None.

        This is how the "Plate" and "Experiment" headers are spotted -
        they are alone on their row, with nothing beside them.
        """
        values = [cell for cell in self.rows.get(row, []) if cell.value is not None]

        return self.text(values[0]) if len(values) == 1 else None

    def field(self, first: int, last: int, label: str) -> str | None:
        """
        Value written beside a label in rows `first` to `last`. None if absent.

        A field is a row that starts with the label and has its value
        somewhere to the right, not necessarily in the next column, so
        the first non-empty cell after the label wins. The first matching
        row wins too; a second one with the same label is ignored.
        """
        wanted = normalize(label)

        for row in range(first, last + 1):
            cells = self.rows.get(row, [])

            # the label only counts as one if it opens the row
            if cells and normalize(self.text(cells[0])) == wanted:
                found = (
                    self.text(cell) for cell in cells[1:] if cell.value is not None
                )

                # no cell after the label, or a blank one, reads as absent
                return next(found, "") or None

        return None

    def has_wells(self, first: int, last: int) -> bool:
        """
        Whether rows `first` - `last` hold a well grid.

        True when some row has a label cell ("A", "B" etc.) with
        at least two numbers to its right.

        This is only meant to tell a real block from an empty one,
        it does not check that the grid is well formed.
        """
        for row in range(first, last + 1):
            cells = self.rows.get(row, [])
            label = next((c for c in cells if ROW_LABEL.match(self.text(c))), None)

            if label is None:
                continue

            # count by column number, because the cells do not need to be adjacent
            numbers = sum(
                1
                for cell in cells
                if cell.col > label.col and isinstance(cell.value, (int, float))
            )

            if numbers >= 2:
                return True

        return False


@dataclass(frozen=True)
class Plate:
    """
    One plate block, and everything needed to write it out again.

    Attributes:
        source: File the plate was read from.
        sheet: Worksheet holding the block.
        styles: Styles of the source file, by ID.
        first_row: The "Plate" header row.
        last_row: Last row of the block.
        preamble: Rows above the first block, which describe the whole
            file rather than this plate. None if there are none.
        name: The block's "Plate name" field, None if it has none.
        experiment: Name of the experiment this plate belongs to.
        read_time: The block's "Read Time" field.
    """

    source: Path
    sheet: Sheet
    styles: dict[str, Style]
    first_row: int
    last_row: int
    preamble: tuple[int, int] | None
    name: str | None
    experiment: str | None
    read_time: str | None

    @property
    def wells(self) -> dict[tuple[int, int], float]:
        """
        The block's numbers, keyed by (row within the block, column).

        Only used to tell whether two plates of the same name really hold
        the same values.
        """
        return {
            (cell.row - self.first_row, cell.col): float(cell.value)
            for row in range(self.first_row, self.last_row + 1)
            for cell in self.sheet.rows.get(row, [])
            if isinstance(cell.value, (int, float))
        }


def read_style(style: Element) -> Style:
    """
    Translate one `<Style>` element.

    Covers what these XML files actually contain: font name, bold and
    colour, alignment, and the number format. Borders, fills and style
    inheritance are ignored because nothing uses them.
    """
    fonts = children(style, "Font")
    alignments = children(style, "Alignment")
    formats = children(style, "NumberFormat")
    colour = attr(fonts[0], "Color") if fonts else None
    code = attr(formats[0], "Format") if formats else None

    return Style(
        font=Font(
            name=attr(fonts[0], "FontName"),
            bold=attr(fonts[0], "Bold") == "1",
            # the XML writes #RRGGBB, openpyxl wants AARRGGBB
            color=f"FF{colour.lstrip('#').upper()}" if colour else None,
        )
        if fonts
        else None,
        alignment=Alignment(
            horizontal=(attr(alignments[0], "Horizontal") or "").lower() or None,
            vertical=(attr(alignments[0], "Vertical") or "").lower() or None,
            wrap_text=attr(alignments[0], "WrapText") == "1",
        )
        if alignments
        else None,
        number_format=NUMBER_FORMATS.get(code.lower(), code) if code else None,
    )


def read_value(cell: Element) -> object:
    """
    Read a cell's `<Data>` as an int, float, or string. None if empty.

    A cell is a number only if it is marked as one. Everything else stays
    a string, so "OVRFLW" and the like survive untouched.
    """
    data = children(cell, "Data")
    # itertext, because inline markup can split the text into several nodes
    text = "".join(data[0].itertext()).strip() if data else ""

    if not text:
        return None

    if attr(data[0], "Type") != "Number":
        return text

    # marked as a number but unparseable: keep it as written
    try:
        value = float(text)
    except ValueError:
        return text

    return int(value) if value.is_integer() else value


def read_sheet(worksheet: Element) -> Sheet:
    """
    Read a `<Worksheet>` into a grid.

    Walks its `<Table>`: the `<Column>` elements carry widths, the
    `<Row>` elements the cells. Positions come from `ss:Index` where the
    file states one and from counting where it does not.
    """
    rows: dict[int, list[Cell]] = {}
    widths: dict[int, float] = {}
    heights: dict[int, float] = {}
    row_index = 0

    tables = children(worksheet, "Table")
    if not tables:
        return Sheet(rows, widths, heights, 0)

    # widths first - these come before the rows and describe whole columns
    column = 0
    for element in children(tables[0], "Column"):
        column = index_of(attr(element, "Index"), column + 1)
        width = attr(element, "Width")
        if width is not None:
            widths[column] = float(width)

    for row in children(tables[0], "Row"):
        row_index = index_of(attr(row, "Index"), row_index + 1)
        height = attr(row, "Height")

        if height is not None:
            heights[row_index] = float(height)

        column = 0
        for element in children(row, "Cell"):
            column = index_of(attr(element, "Index"), column + 1)
            span = index_of(attr(element, "MergeAcross"), 0)

            rows.setdefault(row_index, []).append(
                Cell(
                    row=row_index,
                    col=column,
                    value=read_value(element),
                    style_id=attr(element, "StyleID"),
                    merge_across=span,
                )
            )

            # a merged cell swallows the columns it spans, so the next cell starts past them
            column += span

    return Sheet(rows, widths, heights, row_index)


def read_plates(source: Path) -> list[Plate]:
    """
    Split one XML file into its plate blocks.

    Every worksheet is scanned for header rows: a row holding nothing but
    "Plate" or "Experiment", optionally followed by a bracketed count
    such as "(1 of 4)" or "(3 of 12)". A header owns the rows from itself
    down to the row before the next header, or to the end of the
    worksheet.

    "Plate" blocks become plates. An "Experiment" block yields no plate:
    its "Experiment name" is carried onto every plate below it in that
    worksheet, until the next "Experiment" block. Rows above the first
    header describe the file rather than any one plate and are recorded
    on every plate of the worksheet as its preamble.

    Blocks are not checked for well data here, so a "Plate" header with
    nothing under it still comes back as a plate.

    Returns:
        The plates, by worksheet and then by row. Empty if the file holds
        no "Plate" header at all.

    Raises:
        ValueError: If the file is not well-formed XML.
        OSError: If the file cannot be read.
    """

    try:
        root = parse(source).getroot()
    except ParseError as exc:
        raise ValueError(f"not well-formed XML: {exc}") from exc

    # styles are file-level, shared by every worksheet in it
    styles = {
        style_id: read_style(style)
        for block in root.iter(tag("Styles"))
        for style in children(block, "Style")
        if (style_id := attr(style, "ID")) is not None
    }

    plates: list[Plate] = []

    for worksheet in root.iter(tag("Worksheet")):
        sheet = read_sheet(worksheet)

        # (row, is_plate) for every header, in order down the sheet
        headers: list[tuple[int, bool]] = []

        for row in sorted(sheet.rows):
            label = sheet.only_value(row)
            if label is None:
                continue

            label = normalize(label)
            if PLATE_HEADER.match(label):
                headers.append((row, True))
            elif EXPERIMENT.match(label):
                headers.append((row, False))

        above = [row for row in sorted(sheet.rows) if headers and row < headers[0][0]]
        preamble = (1, above[-1]) if above else None
        experiment: str | None = None

        for position, (first, is_plate) in enumerate(headers):
            # the block runs up to the next header, or to the end
            following = [row for row, _ in headers[position + 1 :]]
            last = following[0] - 1 if following else sheet.last_row

            # an experiment block only names the plates that follow it
            if not is_plate:
                experiment = sheet.field(first, last, "Experiment name")
                continue

            plates.append(
                Plate(
                    source=source,
                    sheet=sheet,
                    styles=styles,
                    first_row=first,
                    last_row=last,
                    preamble=preamble,
                    name=sheet.field(first, last, "Plate name"),
                    experiment=experiment,
                    read_time=sheet.field(first, last, "Read Time"),
                )
            )
    return plates


def sheet_title(wanted: str, used: set[str]) -> str:
    """
    Make a unique, Excel-legal sheet name and record it as taken.

    Excel forbids a handful of characters, cuts names off at 31
    characters and refuses two that differ only in case, so a clash gets
    a " (2)" suffix and the name is trimmed to make room for it.
    """
    cleaned = BAD_SHEET_CHARS.sub("-", wanted).strip("'").strip() or "Plate"
    title = cleaned[:31]
    counter = 2

    while title.casefold() in used:
        suffix = f" ({counter})"
        title = f"{cleaned[: 31 - len(suffix)]}{suffix}"
        counter += 1

    used.add(title.casefold())

    return title


def copy_rows(target: Worksheet, plate: Plate, first: int, last: int, at: int) -> int:
    """
    Copy source rows `first` - `last` to the sheet at row `at`. Return the next free row.

    Rows move up or down as a block. Columns stay where they were, so the
    copy lines up with the original. Styles, merges and row heights come along with it.
    """
    # how far every row shifts, same for all of them
    offset = first - at

    for row in range(first, last + 1):
        for cell in plate.sheet.rows.get(row, []):
            # openpyxl creates the cell on first access
            written = target.cell(row=cell.row - offset, column=cell.col)
            written.value = cell.value
            style = plate.styles.get(cell.style_id or "Default")

            if style is not None:
                style.apply_to(written)

            # merges are a property of the sheet, not of the cell
            if cell.merge_across:
                target.merge_cells(
                    start_row=written.row,
                    start_column=cell.col,
                    end_row=written.row,
                    end_column=cell.col + cell.merge_across,
                )

        height = plate.sheet.heights.get(row)

        if height is not None:
            target.row_dimensions[row - offset].height = height

    return last - offset + 1


def write_plate(target: Worksheet, plate: Plate) -> None:
    """
    Write one plate: provenance rows, the file preamble, then the block.

    The provenance rows are ours, not the file's. They say where the
    plate came from, since a sheet of the merged workbook has otherwise
    lost that. Everything below them is a verbatim copy.
    """
    # styling for the provenance rows is borrowed from the first preamble
    # row, so they do not stand out against the copied content
    template = plate.sheet.rows.get(plate.preamble[0], []) if plate.preamble else []
    provenance = [("Source file", plate.source.name)]
    if plate.experiment is not None:
        provenance.append(("Experiment name", plate.experiment))

    row = 1
    for label, value in provenance:
        for column, text in ((1, label), (2, value)):
            cell = target.cell(row=row, column=column)
            cell.value = text
            source = next((one for one in template if one.col == column), None)
            style = plate.styles.get((source.style_id if source else None) or "Default")

            if style is not None:
                style.apply_to(cell)

        row += 1

    if plate.preamble is not None:
        row = copy_rows(target, plate, *plate.preamble, row)

    # row + 1 leaves a blank row between the preamble and the block
    copy_rows(target, plate, plate.first_row, plate.last_row, row + 1)

    # widths are per column and apply to the whole sheet, so they are set
    # once at the end rather than per copied row
    for column, width in plate.sheet.widths.items():
        letter = get_column_letter(column)
        target.column_dimensions[letter].width = round(width / POINTS_PER_CHARACTER, 2)


def duplicate_status(first: Plate, other: Plate) -> str:
    """
    Describe a skipped duplicate, comparing its wells with the kept plate's.

    Two plates of the same name are normally the same read exported
    twice. If the numbers disagree they are not, and we report it.
    """
    left, right = first.wells, other.wells
    # a position missing from one side counts as differing
    differing = sum(
        1
        for position in left.keys() | right.keys()
        if left.get(position) != right.get(position)
    )

    if differing:
        return (
            f"warning: {differing} well(s) differ from {first.source.name}, kept first"
        )

    return f"skipped: identical to {first.source.name}"


def convert(sources: list[Path], destination: Path) -> list[tuple[str, ...]]:
    """
    Merge plates from several XML files into one workbook.

    Plates are processed in argument order, then source order. A plate whose
    name is already taken is skipped, with a warning if its well values differ
    from the one that was kept. The workbook is saved only if it has a sheet.

    Returns:
        One row per plate: sheet, plate, experiment, read time, source, status.

    Raises:
        OSError: If the workbook cannot be written, for instance because
            the destination is already open in Excel.
    """
    workbook = Workbook()
    # a new workbook opens with one sheet already in it; we remove it
    workbook.remove(workbook.active)
    report: list[tuple[str, ...]] = []
    # first plate seen under each name, for the duplicate comparison
    kept: dict[str, Plate] = {}
    titles: set[str] = set()
    unknown = 0

    for source in sources:
        try:
            plates = read_plates(source)
        except (ValueError, OSError) as exc:
            report.append(("-", "-", "-", "-", source.name, f"error: {exc}"))
            continue

        for plate in plates:
            time = (
                "-"
                if not plate.read_time or plate.read_time.startswith(NEVER_READ)
                else plate.read_time
            )
            found = (plate.name or "-", plate.experiment or "-", time, source.name)
            key = normalize(plate.name) if plate.name else None

            if not plate.sheet.has_wells(plate.first_row, plate.last_row):
                report.append(("-", *found, "skipped: no well data"))
            elif key is not None and key in kept:
                report.append(("-", *found, duplicate_status(kept[key], plate)))
            else:
                # nameless plates are all distinct, so they are numbered
                # rather than deduplicated
                if key is None:
                    unknown += 1
                    wanted = f"Unknown plate #{unknown}"
                else:
                    wanted, kept[key] = plate.name or "", plate
                title = sheet_title(wanted, titles)
                write_plate(workbook.create_sheet(title), plate)
                report.append((title, *found, "written"))

    if workbook.sheetnames:
        workbook.save(destination)

    return report


def format_table(report: list[tuple[str, ...]]) -> list[str]:
    """The report as lines of a padded table, header row first."""
    table = [REPORT_HEADERS, *report]
    # pad each column to its widest entry, header included
    widths = [
        max(len(row[column]) for row in table) for column in range(len(REPORT_HEADERS))
    ]

    return [
        "  ".join(text.ljust(width) for text, width in zip(row, widths)).rstrip()
        for row in table
    ]


@dataclass(frozen=True)
class Summary:
    """
    What a run came to, counted from its report.

    Attributes:
        written: Plates written to a sheet.
        skipped: Plates left out, whether empty or duplicated.
        differing: Duplicates whose values did not match.
        errors: Files that could not be read.
        files: Files that were read.
    """

    written: int
    skipped: int
    differing: int
    errors: int
    files: int

    @classmethod
    def of(cls, report: list[tuple[str, ...]]) -> "Summary":
        """Count up one report."""
        statuses = [row[-1] for row in report]
        written = statuses.count("written")
        errors = sum(1 for status in statuses if status.startswith("error"))

        return cls(
            written=written,
            skipped=len(statuses) - written - errors,
            differing=sum(1 for status in statuses if status.startswith("warning")),
            errors=errors,
            # unreadable files contribute a row each but no plates
            files=len({row[-2] for row in report}) - errors,
        )

    def text(self, destination: Path) -> str:
        """The one-line summary shown after a run."""
        if self.written:
            summary = (
                f"wrote {destination}: {self.written} sheet(s) "
                f"from {self.files} file(s)"
            )
        else:
            summary = "nothing written: no plate with well data was found"
        if self.skipped:
            summary += f", {self.skipped} plate(s) skipped"
        if self.differing:
            summary += f" ({self.differing} with differing values)"
        if self.errors:
            summary += f", {self.errors} file(s) unreadable"

        return summary
